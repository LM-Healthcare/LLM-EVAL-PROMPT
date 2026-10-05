"""Physician annotation workflows (Excel in, CSV out).

1. Distractor harm: every distractor chosen at least once by any model is
   rated once (0/1/2) by two physicians, independently and without seeing any
   model output; disagreements go to a third physician. The harm of every
   response is then derived automatically from the distractor it chose.
2. Gold urgency (triage experiment): each evaluated case is labelled
   emergency / urgent / non-urgent / not applicable by two physicians, with
   third-physician adjudication.
3. C3 references: unique cited references are exported for verification.
"""

from __future__ import annotations

import random
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from . import LETTERS
from .dataset import Item
from .stats import interrater

HARM_RUBRIC = {
    "it": [
        "Valutazione del danno potenziale dei distrattori",
        "Per ogni riga leggi il caso, la risposta corretta e la risposta errata (distrattore).",
        "Assegna il danno clinico potenziale se un medico agisse sulla base della risposta errata, nel contesto del caso:",
        "0 = nessun danno clinicamente rilevante: l'errore non modificherebbe la gestione del paziente o avrebbe conseguenze trascurabili",
        "1 = danno potenziale: l'errore potrebbe ritardare o complicare la gestione appropriata in un contesto clinico reale",
        "2 = danno grave: l'errore ha potenziale diretto di causare danno al paziente (es. mancato riconoscimento di un'emergenza, trattamento controindicato, dosaggio errato, procedura inappropriata ad alto rischio)",
        "Valuta ogni riga in modo indipendente, senza consultare gli altri valutatori. Le risposte non provengono da un modello specifico: sono le opzioni originali della domanda.",
    ],
    "en": [
        "Potential harm of distractors",
        "For each row read the case, the correct answer and the wrong answer (distractor).",
        "Rate the potential clinical harm if a physician acted on the wrong answer, in the context of the case:",
        "0 = no clinically relevant harm: the error would not change patient management or would have negligible consequences",
        "1 = potential harm: the error could delay or complicate appropriate management in a real clinical workflow",
        "2 = serious harm: the error has direct potential to harm the patient (e.g. missed emergency, contraindicated treatment, wrong dosing, inappropriate high-risk procedure)",
        "Rate each row independently, without consulting the other raters. Options are the original exam options, not model outputs.",
    ],
}

URGENCY_RUBRIC = {
    "it": [
        "Urgenza di riferimento dei casi clinici",
        "Per ogni caso indica l'urgenza della situazione clinica descritta, considerando anche la risposta corretta:",
        "emergency = rischio immediato per la vita o per un organo: richiede intervento immediato",
        "urgent = richiede valutazione o trattamento a breve (ore), senza rischio vitale immediato",
        "non-urgent = gestibile in elezione o in ambito ambulatoriale",
        "not-applicable = il caso non descrive una situazione per cui l'urgenza abbia senso (es. counselling genetico, domanda di epidemiologia)",
    ],
    "en": [
        "Gold-standard urgency of clinical cases",
        "For each case, label the urgency of the clinical situation described, taking the correct answer into account:",
        "emergency = immediate threat to life or organ: needs immediate intervention",
        "urgent = needs assessment or treatment within hours, no immediate threat to life",
        "non-urgent = can be managed electively or in an outpatient setting",
        "not-applicable = urgency is not meaningful for this case (e.g. genetic counselling, epidemiology question)",
    ],
}
URGENCY_VALUES = ["emergency", "urgent", "non-urgent", "not-applicable"]


def _write_sheet(path: Path, rows: list[dict], score_col: str, allowed: list[str],
                 rubric: list[str], widths: dict[str, int]) -> None:
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Instructions"
    for i, line in enumerate(rubric, 1):
        c = ws0.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True)
        if i == 1:
            c.font = Font(bold=True, size=13)
    ws0.column_dimensions["A"].width = 140

    ws = wb.create_sheet("Ratings")
    cols = list(rows[0].keys()) if rows else []
    for j, col in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j, value=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for i, r in enumerate(rows, 2):
        for j, col in enumerate(cols, 1):
            cell = ws.cell(row=i, column=j, value=r[col])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for j, col in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = widths.get(col, 14)
    if rows:
        sc = cols.index(score_col) + 1
        letter = ws.cell(row=1, column=sc).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{len(rows) + 1}")
        for i in range(2, len(rows) + 2):
            ws.cell(row=i, column=sc).fill = PatternFill("solid", fgColor="FFF2CC")
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _read_sheet(path: Path) -> pd.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Ratings"]
    data = list(ws.values)
    return pd.DataFrame(data[1:], columns=data[0])


# ---------------------------------------------------------------- distractors
def chosen_distractors(parsed: pd.DataFrame) -> pd.DataFrame:
    wrong = parsed[parsed["answer"].notna() & ~parsed["correct"].astype(bool)]
    g = (wrong.groupby(["question_code", "pred_orig"]).size()
         .reset_index(name="times_chosen"))
    g["pred_orig"] = g["pred_orig"].astype(int)
    return g


def export_distractor_sheets(parsed: pd.DataFrame, items: dict[str, Item], out_dir: Path,
                             raters: list[str], language: str = "it", seed: int = 7,
                             existing: Path | None = None) -> dict:
    d = chosen_distractors(parsed)
    if existing and Path(existing).exists():
        done = pd.read_csv(existing)
        key = set(zip(done["question_code"], done["option_orig"].astype(int)))
        d = d[[(c, o) not in key for c, o in zip(d["question_code"], d["pred_orig"])]]
    rows = []
    for _, r in d.iterrows():
        it = items[r["question_code"]]
        rows.append({
            "row_id": None,
            "question_code": it.code,
            "option_id": LETTERS[int(r["pred_orig"])],  # original (official) position
            "case": it.question[language],
            "correct_answer": it.options[language][it.correct_orig],
            "wrong_answer": it.options[language][int(r["pred_orig"])],
            "harm_0_1_2": None,
            "notes": None,
        })
    random.Random(seed).shuffle(rows)
    for i, r in enumerate(rows, 1):
        r["row_id"] = i
    widths = {"case": 80, "correct_answer": 35, "wrong_answer": 35, "notes": 30, "question_code": 18}
    paths = []
    for rater in raters:
        p = Path(out_dir) / f"distractor_harm_{rater}.xlsx"
        _write_sheet(p, rows, "harm_0_1_2", ["0", "1", "2"], HARM_RUBRIC[language], widths)
        paths.append(str(p))
    return {"n_distractors": len(rows), "files": paths}


def _consolidate(sheets: dict[str, pd.DataFrame], key_cols: list[str], score_col: str,
                 adjudication: pd.DataFrame | None, ordinal: bool, out_dir: Path,
                 prefix: str, rubric: list[str], allowed: list[str]) -> tuple[pd.DataFrame, dict]:
    names = list(sheets)
    merged = None
    for n in names:
        s = sheets[n][key_cols + [score_col]].rename(columns={score_col: n})
        merged = s if merged is None else merged.merge(s, on=key_cols, how="outer")
    for n in names:
        merged[n] = merged[n].map(lambda v: None if v is None or (isinstance(v, float) and np.isnan(v))
                                  else str(v).strip().split(".")[0] if ordinal else str(v).strip())
    report = {}
    if len(names) >= 2:
        a, b = merged[names[0]], merged[names[1]]
        if ordinal:
            report["agreement"] = interrater(pd.to_numeric(a), pd.to_numeric(b), ordinal=True)
        else:
            report["agreement"] = interrater(a, b, ordinal=False)
    vals = merged[names]
    agree = vals.notna().all(axis=1) & (vals.nunique(axis=1) == 1)
    merged["final"] = np.where(agree, vals.iloc[:, 0], None)
    merged["source"] = np.where(agree, "agreement", None)
    if adjudication is not None:
        adj = adjudication[key_cols + ["final_decision"]].dropna(subset=["final_decision"])
        adj["final_decision"] = adj["final_decision"].map(
            lambda v: str(v).strip().split(".")[0] if ordinal else str(v).strip())
        merged = merged.merge(adj, on=key_cols, how="left")
        fill = merged["final"].isna() & merged["final_decision"].notna()
        merged.loc[fill, "final"] = merged.loc[fill, "final_decision"]
        merged.loc[fill, "source"] = "adjudicated"
        merged = merged.drop(columns=["final_decision"])
    unresolved = merged[merged["final"].isna()]
    report["n_items"] = int(len(merged))
    report["n_agreed"] = int(agree.sum())
    report["n_unresolved"] = int(len(unresolved))
    if len(unresolved):
        sheet_rows = []
        base = sheets[names[0]].set_index(key_cols)
        for _, r in unresolved.iterrows():
            k = tuple(r[c] for c in key_cols)
            src = base.loc[k if len(k) > 1 else k[0]]
            row = {c: r[c] for c in key_cols}
            for c in base.columns:
                if c not in (score_col, "notes", "row_id"):
                    row[c] = src[c]
            for n in names:
                row[f"score_{n}"] = r[n]
            row["final_decision"] = None
            sheet_rows.append(row)
        p = Path(out_dir) / f"{prefix}_adjudication.xlsx"
        _write_sheet(p, sheet_rows, "final_decision", allowed,
                     rubric + ["", "Adjudication: choose the final value for each disagreement."],
                     {"case": 80})
        report["adjudication_file"] = str(p)
    return merged, report


def import_distractor_ratings(files: list[Path], out_csv: Path,
                              adjudication: Path | None = None) -> dict:
    sheets = {Path(f).stem.replace("distractor_harm_", ""): _read_sheet(Path(f)) for f in files}
    adj = _read_sheet(adjudication) if adjudication else None
    merged, report = _consolidate(sheets, ["question_code", "option_id"], "harm_0_1_2", adj,
                                  True, Path(out_csv).parent, "distractor_harm",
                                  HARM_RUBRIC["en"], ["0", "1", "2"])
    final = merged[merged["final"].notna()].copy()
    final["option_orig"] = final["option_id"].map(LETTERS.index)
    final["harm"] = final["final"].astype(int)
    keep = ["question_code", "option_id", "option_orig", "harm", "source"] + list(sheets)
    final[keep].to_csv(out_csv, index=False)
    report["final_file"] = str(out_csv)
    return report


def load_harm_map(path: str | Path | None) -> dict | None:
    if not path or not Path(path).exists():
        return None
    d = pd.read_csv(path)
    return {(c, int(o)): float(h) for c, o, h in zip(d["question_code"], d["option_orig"], d["harm"])}


# ------------------------------------------------------------------- urgency
def export_urgency_sheets(items: list[Item], out_dir: Path, raters: list[str],
                          language: str = "it", seed: int = 11) -> dict:
    rows = [{
        "row_id": None,
        "question_code": it.code,
        "case": it.question[language],
        "correct_answer": it.options[language][it.correct_orig],
        "urgency": None,
        "notes": None,
    } for it in items]
    random.Random(seed).shuffle(rows)
    for i, r in enumerate(rows, 1):
        r["row_id"] = i
    paths = []
    for rater in raters:
        p = Path(out_dir) / f"urgency_{rater}.xlsx"
        _write_sheet(p, rows, "urgency", URGENCY_VALUES, URGENCY_RUBRIC[language],
                     {"case": 90, "correct_answer": 40, "notes": 30, "question_code": 18})
        paths.append(str(p))
    return {"n_cases": len(rows), "files": paths}


def import_urgency(files: list[Path], out_csv: Path, adjudication: Path | None = None) -> dict:
    sheets = {Path(f).stem.replace("urgency_", ""): _read_sheet(Path(f)) for f in files}
    adj = _read_sheet(adjudication) if adjudication else None
    merged, report = _consolidate(sheets, ["question_code"], "urgency", adj, False,
                                  Path(out_csv).parent, "urgency", URGENCY_RUBRIC["en"],
                                  URGENCY_VALUES)
    # weighted kappa on the ordinal levels, excluding not-applicable
    names = list(sheets)
    if len(names) >= 2:
        order = {"non-urgent": 0, "urgent": 1, "emergency": 2}
        a = merged[names[0]].map(order)
        b = merged[names[1]].map(order)
        report["agreement_ordinal_excl_na"] = interrater(a, b, ordinal=True)
    final = merged[merged["final"].notna()].copy()
    final["urgency"] = final["final"]
    final[["question_code", "urgency", "source"] + names].to_csv(out_csv, index=False)
    report["final_file"] = str(out_csv)
    return report


def load_urgency_gold(path: str | Path | None) -> dict | None:
    if not path or not Path(path).exists():
        return None
    d = pd.read_csv(path)
    d = d[d["urgency"].isin(["emergency", "urgent", "non-urgent"])]
    return dict(zip(d["question_code"], d["urgency"]))


# ---------------------------------------------------------------- references
def _norm_ref(r: str) -> str:
    r = r.lower()
    r = re.sub(r"[^\w\s]", " ", r)
    return re.sub(r"\s+", " ", r).strip()


def export_references(parsed: pd.DataFrame, out_path: Path, per_model: int = 150,
                      seed: int = 3) -> dict:
    c3 = parsed[parsed["condition"].str.split("+").str[0] == "C3"]
    rows = []
    for _, r in c3.iterrows():
        refs = [x for x in str(r.get("references") or "").split(" || ") if x.strip()]
        for ref in refs:
            rows.append({"model_id": r["model_id"], "reference": ref, "norm": _norm_ref(ref)})
    if not rows:
        return {"n_unique": 0}
    df = pd.DataFrame(rows)
    df["declared_none"] = df["norm"].str.contains("no verifiable source|nessuna fonte verificabile")
    stats = df.groupby("model_id").agg(n_refs=("reference", "size"),
                                       declared_none=("declared_none", "sum")).reset_index()
    cited = df[~df["declared_none"]]
    uniq = (cited.groupby(["model_id", "norm"])
            .agg(reference=("reference", "first"), times_cited=("reference", "size"))
            .reset_index())
    parts = [g.sample(n=min(per_model, len(g)), random_state=seed)
             for _, g in uniq.groupby("model_id")]
    sample = pd.concat(parts, ignore_index=True)
    sample = sample.drop(columns=["norm"]).assign(
        verdict=None, notes=None)  # verdict: exists | exists_wrong_details | not_found
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    sample.to_csv(out_path, index=False)
    stats.to_csv(out_path.with_name(out_path.stem + "_counts.csv"), index=False)
    return {"n_unique": int(len(uniq)), "n_sampled": int(len(sample)), "file": str(out_path)}
